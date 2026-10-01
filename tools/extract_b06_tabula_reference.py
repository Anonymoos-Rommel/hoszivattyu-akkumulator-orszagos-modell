"""Pinned source extraction only; no Excel execution or recalculation."""
import argparse
from pathlib import Path
import hashlib,json,datetime
import openpyxl
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--xlsx',type=Path,required=True)
parser.add_argument('--output',type=Path,required=True)
args=parser.parse_args()
SOURCE=args.xlsx
manifest=json.loads((ROOT/'registry/b06_tabula_reference_manifest.json').read_text())
SHA=manifest['source_sha256']
if hashlib.sha256(SOURCE.read_bytes()).hexdigest()!=SHA:raise ValueError('source revision mismatch')
keys=set(['A_C_Ref', 'A_Calc_Door_1', 'A_Calc_Floor_1', 'A_Calc_Floor_2', 'A_Calc_Roof_1', 'A_Calc_Roof_2', 'A_Calc_Wall_1', 'A_Calc_Wall_2', 'A_Calc_Wall_3', 'A_Calc_Window_1', 'A_Calc_Window_2', 'A_Calc_Window_East', 'A_Calc_Window_Horizontal', 'A_Calc_Window_North', 'A_Calc_Window_South', 'A_Calc_Window_West', 'Check_EnvArea_ExactToEstim', 'Check_EnvSum_ExactToEstim', 'Check_FloorArea_ExactToEstim', 'Check_ToBeApplied_FloorArea_ExactToEstim', 'Check_WindowArea_ExactToEstim', 'Code_Building', 'Code_BuildingVariant', 'Code_MeasureType_Door_1', 'Code_MeasureType_Floor_1', 'Code_MeasureType_Floor_2', 'Code_MeasureType_Roof_1', 'Code_MeasureType_Roof_2', 'Code_MeasureType_Wall_1', 'Code_MeasureType_Wall_2', 'Code_MeasureType_Wall_3', 'Code_MeasureType_Window_1', 'Code_MeasureType_Window_2', 'Code_Measure_Window_1', 'Code_Measure_Window_2', 'Code_StatusDataset', 'Code_ThermalBridging_Refurbished', 'Code_TypeVariant', 'Date_Change', 'Date_Entry', 'Description_BuildingVariant', 'F_f', 'F_red_htr1', 'F_red_htr4', 'F_red_temp', 'F_sh_hor', 'F_sh_vert', 'F_w', 'HeatingDays', 'I_Sol_East', 'I_Sol_Hor', 'I_Sol_North', 'I_Sol_South', 'I_Sol_West', 'R_Add_UnheatedSpace_Floor_1', 'R_Add_UnheatedSpace_Floor_2', 'R_Add_UnheatedSpace_Roof_1', 'R_Add_UnheatedSpace_Roof_2', 'R_Add_UnheatedSpace_Wall_1', 'R_Add_UnheatedSpace_Wall_2', 'R_Add_UnheatedSpace_Wall_3', 'R_Measure_Door_1', 'R_Measure_Floor_1', 'R_Measure_Floor_2', 'R_Measure_Roof_1', 'R_Measure_Roof_2', 'R_Measure_Wall_1', 'R_Measure_Wall_2', 'R_Measure_Wall_3', 'R_Measure_Window_1', 'R_Measure_Window_2', 'Remark_BuildingVariant_1', 'Remark_BuildingVariant_2', 'Sum_DeltaT_for_HeatingDays', 'Theta_e', 'U_Door_1', 'U_Floor_1', 'U_Floor_2', 'U_Roof_1', 'U_Roof_2', 'U_Wall_1', 'U_Wall_2', 'U_Wall_3', 'U_Window_1', 'U_Window_2', 'b_Transmission_Floor_1', 'b_Transmission_Floor_2', 'b_Transmission_Roof_1', 'b_Transmission_Roof_2', 'b_Transmission_Wall_1', 'b_Transmission_Wall_2', 'b_Transmission_Wall_3', 'c_m', 'd_Insulation_Floor_1', 'd_Insulation_Floor_2', 'd_Insulation_Roof_1', 'd_Insulation_Roof_2', 'd_Insulation_Wall_1', 'd_Insulation_Wall_2', 'd_Insulation_Wall_3', 'delta_U_ThermalBridging_Original', 'delta_U_ThermalBridging_Refurbished', 'f_Measure_Door_1', 'f_Measure_Floor_1', 'f_Measure_Floor_2', 'f_Measure_Roof_1', 'f_Measure_Roof_2', 'f_Measure_Wall_1', 'f_Measure_Wall_2', 'f_Measure_Wall_3', 'f_Measure_Window_1', 'f_Measure_Window_2', 'g_gl_n_Measure_Window_1', 'g_gl_n_Measure_Window_2', 'g_gl_n_Window_1', 'g_gl_n_Window_2', 'h_room', 'n_Apartment', 'n_air_infiltration', 'n_air_use', 'phi_int', 'q_h_nd', 'r_EnvFloor_ExactToEstim', 'r_EnvTotal_ExactToEstim', 'r_EnvWindow_ExactToEstim', 'theta_i'])
cached=openpyxl.load_workbook(SOURCE,read_only=True,data_only=True)
formulas=openpyxl.load_workbook(SOURCE,read_only=True,data_only=False)
s=cached['Calc.Set.Building'];f=formulas[s.title]
header=next(s.iter_rows(min_row=1,max_row=1,values_only=True))
if not keys.issubset(header):raise ValueError('missing source columns')
columns={k:header.index(k)+1 for k in sorted(keys)}
units=next(s.iter_rows(min_row=4,max_row=4,values_only=True))
selected={1218,1219,1220,1227,1228,1229,1239,1240,1241}
rows=[]
for number,(values,raw) in enumerate(zip(s.iter_rows(min_row=1218,max_row=1241,values_only=True),f.iter_rows(min_row=1218,max_row=1241,values_only=True)),1218):
 if number not in selected:continue
 vals={k:values[c-1] for k,c in columns.items()}
 vals={k:str(v) if isinstance(v,(datetime.datetime,datetime.date)) else v for k,v in vals.items()}
 rows.append({'source_row':number,'values':vals,'formula_cells':{k:raw[c-1] for k,c in columns.items() if isinstance(raw[c-1],str) and raw[c-1].startswith('=')}})
result=dict(source_id='SRC-B06-TABULA-HU-CALCULATOR-2016',source_sha256=SHA,sheet=s.title,attribution='IEE Projects TABULA + EPISCOPE (www.episcope.eu)',scope='HISTORICAL_SEASONAL_BUILDING_REFERENCE_NOT_NATIONAL_CALIBRATION',columns={k:dict(column=get_column_letter(c),native_unit=units[c-1]) for k,c in columns.items()},rows=rows)
args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print({'rows':len(rows),'fields':len(keys),'formula_cells':sum(len(r['formula_cells']) for r in rows),'bytes':args.output.stat().st_size})
